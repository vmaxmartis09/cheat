# -*- coding: utf-8 -*-
import json
import openpyxl

# Load answer key
wb = openpyxl.load_workbook('docs2/extracted/Session 11_ Final trial test correction (part 1)/Session 11: Final trial test correction (part 1)/Master answer sheet for final test.xlsx', data_only=True)
ws = wb['answer Key for final test']
answers = {}
for r in range(3, 28):
    for c_q, c_a in [(10, 11), (12, 13), (14, 15), (16, 17)]:
        q = ws.cell(r, c_q).value
        a = ws.cell(r, c_a).value
        if q and 147 <= int(q) <= 200:
            answers[int(q)] = a

print("Loaded Part 7 answer count:", len(answers))
