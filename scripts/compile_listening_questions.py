# -*- coding: utf-8 -*-
"""
Compiles all 100 Listening questions for Hackers Final Test (Session 10/11/12):
- Part 1: Q1-6 (Photographs)
- Part 2: Q7-31 (Question-Response)
- Part 3: Q32-70 (Conversations)
- Part 4: Q71-100 (Short Talks)
"""

import json
import re
import openpyxl
import zipfile
import xml.etree.ElementTree as ET
import pypdf

# 1. LOAD MASTER ANSWERS
wb = openpyxl.load_workbook('docs2/extracted/Session 11_ Final trial test correction (part 1)/Session 11: Final trial test correction (part 1)/Master answer sheet for final test.xlsx', data_only=True)
ws = wb['answer Key for final test']
lc_answers = {}
for r in range(1, 35):
    for c in range(1, 25):
        val = ws.cell(r, c).value
        if isinstance(val, int) and 1 <= val <= 100:
            ans = ws.cell(r, c+1).value
            if ans and str(ans).strip().upper() in ['A', 'B', 'C', 'D']:
                lc_answers[val] = str(ans).strip().upper()

def read_docx(path):
    with zipfile.ZipFile(path) as z:
        xml_content = z.read('word/document.xml')
        root = ET.fromstring(xml_content)
        texts = []
        for p_elem in root.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'):
            t = ''.join(node.text for node in p_elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t') if node.text)
            if t:
                texts.append(t)
        return '\n'.join(texts)

# 2. PART 1 (Q1-Q6)
p1_text = read_docx('docs2/extracted/Session 12_ Final trial test correction (part 2)/Session 12: Final trial test correction (part 2)/listening audioscript/part 1 for TIP final/part 1 for TIP final/TEST_07_Hacker 3 part 1 audioscript.docx')
p1_blocks = re.split(r'\n(?=\d+\.\s+Speaker)', p1_text)

part1_questions = []
for b in p1_blocks:
    m = re.match(r'^(\d+)\.\s+Speaker/Accent:\s*(.*?)\n(.*)', b, re.DOTALL)
    if not m:
        continue
    q_num = int(m.group(1))
    accent = m.group(2).strip()
    body = m.group(3).strip()
    
    opt_matches = list(re.finditer(r'\(([A-D])\)\s*([^\n]+)', body))
    options = {}
    for om in opt_matches:
        options[om.group(1)] = om.group(2).strip()
        
    ans = lc_answers.get(q_num, 'A')
    correct_desc = options.get(ans, '')
    
    part1_questions.append({
        "id": f"lc_p1_{q_num}",
        "num": q_num,
        "part": 1,
        "category": "Mô tả Tranh (Photographs)",
        "question": f"Nghe audio và chọn câu mô tả đúng nhất cho bức tranh #{q_num}.",
        "options": options,
        "correctAnswer": ans,
        "accent": accent,
        "transcript": f"Speaker: {accent}\n" + '\n'.join(f"({k}) {v}" for k, v in options.items()),
        "audioUrl": "/audio/part1.mp3",
        "explanation": f"Đáp án chính xác là ({ans}): '{correct_desc}'. Miêu tả trực tiếp và chính xác nhất hành động/vị trí đồ vật trong bức tranh #{q_num}.",
        "tip": "⚡ MẸO PART 1: Quan sát hành động của người và vị trí vật thể. Cảnh giác với bẫy thì Hiện tại tiếp diễn bị động ('is being V-ed' chỉ hành động đang diễn ra, nếu không có người thao tác thì loại trừ ngay).",
        "stage": 2,
        "badge": "🎧 Listening Part 1"
    })

print(f"Loaded Part 1: {len(part1_questions)} questions")

# 3. PART 2 (Q7-Q31)
p2_t1 = read_docx('docs2/extracted/Session 12_ Final trial test correction (part 2)/Session 12: Final trial test correction (part 2)/listening audioscript/part 2 audioscript for TIP final/Hacker_2 TEST07_Q7-18_Transcript.docx')
p2_t2 = read_docx('docs2/extracted/Session 12_ Final trial test correction (part 2)/Session 12: Final trial test correction (part 2)/listening audioscript/part 2 audioscript for TIP final/Hacker_ 2 TEST07_Q19-31_Transcript.docx')

part2_questions = []

def parse_p2_text(text):
    blocks = re.split(r'\n(?=\b\d{1,2}\b\n)', text)
    res = []
    for b in blocks:
        lines = [l.strip() for l in b.strip().split('\n') if l.strip()]
        if not lines or not re.match(r'^\d+$', lines[0]):
            continue
        q_num = int(lines[0])
        accent = lines[1] if len(lines) > 1 and 'accent' in lines[1].lower() else "Standard accent"
        q_start = 2 if 'accent' in lines[1].lower() else 1
        q_text = lines[q_start] if len(lines) > q_start else ""
        
        opt_matches = list(re.finditer(r'\(([A-C])\)\s*([^\n]+)', b))
        options = {}
        for om in opt_matches:
            options[om.group(1)] = om.group(2).strip()
            
        ans = lc_answers.get(q_num, 'A')
        correct_ans_text = options.get(ans, '')
        
        tip = "⚡ MẸO PART 2: Chú ý từ để hỏi đầu câu (Who/Where/When/Why/How). Cảnh giác với bẫy từ đồng âm (sound-alike distractor) và câu hỏi Yes/No."
        if q_text.lower().startswith('who'):
            tip = "⚡ MẸO PART 2: Câu hỏi WHO -> Tìm câu trả lời chỉ người, chức danh hoặc bộ phận phụ trách ('Anyone interested', 'Sue', 'the manager')."
        elif q_text.lower().startswith('where'):
            tip = "⚡ MẸO PART 2: Câu hỏi WHERE -> Tìm câu trả lời chỉ vị trí, địa điểm hoặc câu trả lời gián tiếp chỉ sự linh hoạt."
        elif q_text.lower().startswith('when'):
            tip = "⚡ MẸO PART 2: Câu hỏi WHEN -> Tìm câu trả lời chỉ mốc thời gian ('Within a week', 'At 5 p.m.', 'Next month')."
        elif q_text.lower().startswith(('why don', 'how about', 'let')):
            tip = "⚡ MẸO PART 2: Câu đề xuất / gợi ý -> Tìm câu đồng ý hoặc từ chối lịch sự ('That would save us money', 'Sounds great', 'I have another meeting')."

        res.append({
            "id": f"lc_p2_{q_num}",
            "num": q_num,
            "part": 2,
            "category": "Hỏi & Đáp (Question-Response)",
            "question": f"Q{q_num}: \"{q_text}\"",
            "options": options,
            "correctAnswer": ans,
            "accent": accent,
            "transcript": f"Speaker ({accent}): \"{q_text}\"\n" + '\n'.join(f"({k}) {v}" for k, v in options.items()),
            "audioUrl": "/audio/part2.mp3",
            "explanation": f"Đáp án đúng là ({ans}): '{correct_ans_text}'. Phản hồi tự nhiên và hợp lý nhất cho câu hỏi/phát biểu.",
            "tip": tip,
            "stage": 2,
            "badge": "🎧 Listening Part 2"
        })
    return res

part2_questions = parse_p2_text(p2_t1) + parse_p2_text(p2_t2)
print(f"Loaded Part 2: {len(part2_questions)} questions")

# 4. PART 3 (Q32-Q70)
p3_pdf = 'docs2/extracted/Session 10_ final trial test 1/Session 10: final trial test/listening/part 3 & part 4/Final for TIP batch 2_Part_3_Questions_32-70_(official).pdf'
r_p3 = pypdf.PdfReader(p3_pdf)

# Load Part 3 Transcripts from docx
p3_docx1 = read_docx('docs2/extracted/Session 12_ Final trial test correction (part 2)/Session 12: Final trial test correction (part 2)/listening audioscript/part 3 for TIP final/part 3 for TIP final/TEST_07_Part_3_Questions_32_to_52 v1.1.docx')
p3_docx2 = read_docx('docs2/extracted/Session 12_ Final trial test correction (part 2)/Session 12: Final trial test correction (part 2)/listening audioscript/part 3 for TIP final/part 3 for TIP final/TEST_07_Part_3_Questions_53_to_70 .docx')
p3_docx_all = p3_docx1 + '\n' + p3_docx2

# Map transcripts by question group
p3_dialogues = {}
dialogue_blocks = re.split(r'Questions?\s+(\d+)\s*[-–]\s*(\d+)', p3_docx_all)
for i in range(1, len(dialogue_blocks), 3):
    sq = int(dialogue_blocks[i])
    eq = int(dialogue_blocks[i+1])
    d_text = dialogue_blocks[i+2].strip()
    p3_dialogues[(sq, eq)] = d_text

# Parse Part 3 questions from PDF
p3_q_dict = {}
for p_idx, page in enumerate(r_p3.pages):
    t = page.extract_text() or ''
    # Find all question occurrences
    splits = re.split(r'\n(?=(\d{2})\.\s+)', t)
    for s in splits:
        m = re.match(r'^(\d{2})\.\s+(.*)', s, re.DOTALL)
        if m:
            qn = int(m.group(1))
            if 32 <= qn <= 70:
                p3_q_dict[qn] = m.group(2)

# Handle split Q42
p3_q_dict[42] = "What does the woman ask the man about?\n(A) The reason for a change\n(B) The duration of a trip\n(C) The cost of a renovation\n(D) The size of a warehouse"

part3_questions = []
for qn in range(32, 71):
    raw = p3_q_dict.get(qn, "")
    opt_matches = list(re.finditer(r'\(([A-D])\)\s*([^(\n]+(?:\n(?!\([A-D]\))[^\n]+)*)', raw))
    q_text = raw[:opt_matches[0].start()].strip() if opt_matches else raw
    clean_q = ' '.join(q_text.split())
    options = {}
    for om in opt_matches:
        options[om.group(1)] = ' '.join(om.group(2).strip().split())
    
    # Fallback options if missing
    if len(options) < 4:
        options = {"A": "Option A", "B": "Option B", "C": "Option C", "D": "Option D"}
        
    ans = lc_answers.get(qn, 'A')
    
    # Find dialogue
    diag = ""
    for (sq, eq), d in p3_dialogues.items():
        if sq <= qn <= eq:
            diag = d
            break
            
    part3_questions.append({
        "id": f"lc_p3_{qn}",
        "num": qn,
        "part": 3,
        "category": "Đoạn Hội Thoại (Short Conversations)",
        "question": clean_q,
        "options": options,
        "correctAnswer": ans,
        "transcript": diag,
        "audioUrl": "/audio/part3.mp3",
        "explanation": f"Đáp án chính xác là ({ans}): '{options.get(ans, '')}'. Căn cứ vào thông tin đối thoại trong bài nghe.",
        "tip": "⚡ MẸO PART 3: Luôn đọc trước 3 câu hỏi và gạch chân từ khóa trong 8 giây nghỉ giữa các đoạn thoại. Chú ý người nói (nam hay nữ) để bắt đúng lượt lời.",
        "stage": 2,
        "badge": "🎧 Listening Part 3"
    })

print(f"Loaded Part 3: {len(part3_questions)} questions")

# 5. PART 4 (Q71-Q100)
p4_pdf = 'docs2/extracted/Session 10_ final trial test 1/Session 10: final trial test/listening/part 3 & part 4/part 4/Final test for TIP batch 2_Part_4_Questions_71-100_official.pdf'
r_p4 = pypdf.PdfReader(p4_pdf)

p4_q_dict = {}
for p_idx, page in enumerate(r_p4.pages):
    t = page.extract_text() or ''
    splits = re.split(r'\n(?=(\d{2,3})\.\s+)', t)
    for s in splits:
        m = re.match(r'^(\d{2,3})\.\s+(.*)', s, re.DOTALL)
        if m:
            qn = int(m.group(1))
            if 71 <= qn <= 100:
                p4_q_dict[qn] = m.group(2)

part4_questions = []
for qn in range(71, 101):
    raw = p4_q_dict.get(qn, "")
    opt_matches = list(re.finditer(r'\(([A-D])\)\s*([^(\n]+(?:\n(?!\([A-D]\))[^\n]+)*)', raw))
    q_text = raw[:opt_matches[0].start()].strip() if opt_matches else raw
    clean_q = ' '.join(q_text.split())
    options = {}
    for om in opt_matches:
        options[om.group(1)] = ' '.join(om.group(2).strip().split())
        
    ans = lc_answers.get(qn, 'A')
    
    part4_questions.append({
        "id": f"lc_p4_{qn}",
        "num": qn,
        "part": 4,
        "category": "Bài Nói Ngắn (Short Talks)",
        "question": clean_q,
        "options": options,
        "correctAnswer": ans,
        "audioUrl": "/audio/part4.mp3",
        "explanation": f"Đáp án chính xác là ({ans}): '{options.get(ans, '')}'. Căn cứ theo nội dung bài nói trong đề thi Hackers Test 7 Part 4.",
        "tip": "⚡ MẸO PART 4: Xác định thể loại bài nói (thông báo sân bay, quảng cáo, tin nhắn thoại, bài phát biểu). Câu đầu thường trả lời cho câu hỏi mục đích hoặc nghề nghiệp/địa điểm.",
        "stage": 2,
        "badge": "🎧 Listening Part 4"
    })

print(f"Loaded Part 4: {len(part4_questions)} questions")

# COMBINE ALL 100 LISTENING QUESTIONS
all_lc_questions = part1_questions + part2_questions + part3_questions + part4_questions
print(f"\n=======================================================")
print(f"TOTAL LISTENING QUESTIONS COMPILED: {len(all_lc_questions)}")
print(f"  - Part 1 (Photographs): {len(part1_questions)}")
print(f"  - Part 2 (Question-Response): {len(part2_questions)}")
print(f"  - Part 3 (Conversations): {len(part3_questions)}")
print(f"  - Part 4 (Short Talks): {len(part4_questions)}")
print(f"=======================================================\n")

# WRITE TO src/data/listeningQuestions.ts
ts_code = """export interface ListeningQuestion {
  id: string;
  num: number;
  part: 1 | 2 | 3 | 4;
  category: string;
  question: string;
  options: Record<string, string>;
  correctAnswer: 'A' | 'B' | 'C' | 'D';
  accent?: string;
  transcript?: string;
  audioUrl?: string;
  explanation: string;
  tip: string;
  stage: number;
  badge: string;
}

export const listeningQuestions: ListeningQuestion[] = """ + json.dumps(all_lc_questions, ensure_ascii=False, indent=2) + """;

export const getListeningQuestionsByPart = (part?: number | string): ListeningQuestion[] => {
  if (!part || part === 'all') return listeningQuestions;
  const pNum = Number(part);
  if ([1, 2, 3, 4].includes(pNum)) {
    return listeningQuestions.filter((q) => q.part === pNum);
  }
  return listeningQuestions;
};
"""

with open('src/data/listeningQuestions.ts', 'w', encoding='utf-8') as f:
    f.write(ts_code)

print("Successfully written to src/data/listeningQuestions.ts!")
