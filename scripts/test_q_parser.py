# -*- coding: utf-8 -*-
import re
import openpyxl

wb = openpyxl.load_workbook('docs2/extracted/Session 11_ Final trial test correction (part 1)/Session 11: Final trial test correction (part 1)/Master answer sheet for final test.xlsx', data_only=True)
ws = wb['answer Key for final test']
answers = {}
for r in range(3, 28):
    for c_q, c_a in [(10, 11), (12, 13), (14, 15), (16, 17)]:
        q = ws.cell(r, c_q).value
        a = ws.cell(r, c_a).value
        if q and 147 <= int(q) <= 200:
            answers[int(q)] = a

with open('docs2/extracted/part7_all_ocr.txt') as f:
    text = f.read()

# Clean noise
text = re.sub(r'=== PAGE \d+ ===', '', text)
text = re.sub(r'TEST\s*8\s*PART\s*7\s*\d+', '', text)
text = re.sub(r'GO ON TO THE NEXT PAGE', '', text)
text = re.sub(r'Hackers[Ii]ngang\.com', '', text)

matches = list(re.finditer(r'Questions?\s+(\d+)\s*[-–]\s*(\d+)\s+refer to the following\s+([^\n.]+)', text))

parsed_questions = []

for i, m in enumerate(matches):
    sq, eq = int(m.group(1)), int(m.group(2))
    start_pos = m.start()
    end_pos = matches[i+1].start() if i+1 < len(matches) else len(text)
    block = text[start_pos:end_pos]
    
    # The passage is before the first question (sq.)
    first_q_match = re.search(rf'\b{sq}\.\s+', block)
    if not first_q_match:
        print(f"ERROR: Q{sq} not found in passage {i+1}")
        continue
    
    passage_text = block[:first_q_match.start()].strip()
    questions_block = block[first_q_match.start():].strip()
    
    # Now parse individual questions
    for q_num in range(sq, eq+1):
        if q_num < eq:
            next_q_pat = rf'\b{q_num+1}\.\s+'
            q_match = re.search(rf'\b{q_num}\.\s+(.*?)(?={next_q_pat})', questions_block, re.DOTALL)
        else:
            q_match = re.search(rf'\b{q_num}\.\s+(.*)', questions_block, re.DOTALL)
        
        if not q_match:
            print(f"ERROR: Could not parse Q{q_num}")
            continue
        
        q_raw = q_match.group(1).strip()
        # Find options A, B, C, D
        opt_matches = list(re.finditer(r'\(([A-D])\)\s*([^(\n]+(?:\n(?!\([A-D]\))[^\n]+)*)', q_raw))
        q_text = q_raw[:opt_matches[0].start()].strip() if opt_matches else q_raw
        
        options = {}
        for opt_m in opt_matches:
            opt_letter = opt_m.group(1)
            opt_val = ' '.join(opt_m.group(2).strip().split())
            options[opt_letter] = opt_val
            
        parsed_questions.append({
            'num': q_num,
            'question': ' '.join(q_text.split()),
            'options': options,
            'ans': answers.get(q_num),
            'has_all_opts': len(options) == 4
        })

print(f"Parsed {len(parsed_questions)} questions")
missing_opts = [q['num'] for q in parsed_questions if not q['has_all_opts']]
print("Questions missing some options:", missing_opts)
for q in parsed_questions[:5]:
    print(f"Q{q['num']}: {q['question']}")
    print(f"   Opts: {q['options']}")
    print(f"   Ans: {q['ans']}")
