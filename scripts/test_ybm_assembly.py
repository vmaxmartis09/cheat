# -*- coding: utf-8 -*-
import openpyxl

wb = openpyxl.load_workbook('docs2/extracted/Session 7_ Correction of Part VII assignment/Session 7: Correction of Part VII assignment/master answer sheet for assignment part 7 v1.0.xlsx', data_only=True)
ws = wb['answer Key for test 1 (YBM)']
ybm_answers = {}
for r in range(1, 35):
    for c in range(1, 20):
        val = ws.cell(r, c).value
        if isinstance(val, int) and 101 <= val <= 200:
            ans = ws.cell(r, c+1).value
            if ans and str(ans).strip().upper() in ['A', 'B', 'C', 'D']:
                ybm_answers[val] = str(ans).strip().upper()

print(f"Loaded {len(ybm_answers)} answers for YBM Test 1")
