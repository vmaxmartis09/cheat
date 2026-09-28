# -*- coding: utf-8 -*-
import re
import openpyxl

wb = openpyxl.load_workbook('docs2/extracted/Session 7_ Correction of Part VII assignment/Session 7: Correction of Part VII assignment/master answer sheet for assignment part 7 v1.0.xlsx', data_only=True)
ws = wb['answer Key for test 1 (YBM)']
ybm_answers = {}
for r in range(1, 35):
    for c in range(1, 20):
        val = ws.cell(r, c).value
        if isinstance(val, int) and 101 <= val <= 200:
            ans = ws.cell(r, c+1).value
            if ans and str(ans).strip().upper() in ['A', 'B', 'C', 'D']:
                ybm_answers[val] = str(ans).strip().upper()

with open('docs2/extracted/ybm_part5_ocr.txt') as f:
    text = f.read()

# Clean noise
text = re.sub(r'=== PAGE \d+ ===', '', text)
text = re.sub(r'YBM RC 1000.*', '', text)
text = re.sub(r'GO ON TO THE NEXT PAGE', '', text)

blocks = [b.strip() for b in re.split(r'\n(?=1\d\d\.)', text) if b.strip()]

ybm_p5_list = []
for b in blocks:
    if not re.match(r'^1\d\d\.', b):
        continue
    q_num = int(b[:3])
    # Extract options
    opt_matches = list(re.finditer(r'\(([A-D])\)\s*([^(\n]+(?:\n(?!\([A-D]\))[^\n]+)*)', b))
    if not opt_matches:
        continue
    q_text = b[:opt_matches[0].start()].strip()
    # Remove question number
    q_text = re.sub(r'^\d{3}\.\s*', '', q_text).replace('\n', ' ')
    
    options = {}
    for om in opt_matches:
        letter = om.group(1)
        val = ' '.join(om.group(2).strip().split())
        options[letter] = val
        
    ans = ybm_answers.get(q_num, 'A')
    ybm_p5_list.append({
        'num': q_num,
        'question': q_text,
        'options': options,
        'ans': ans,
        'opt_count': len(options)
    })

print(f"Parsed {len(ybm_p5_list)} YBM Part 5 questions")
missing = [q['num'] for q in ybm_p5_list if q['opt_count'] < 4]
print("Missing opts:", missing)
for q in ybm_p5_list[:5]:
    print(f"Q{q['num']}: {q['question']}")
    print(f"   Opts: {q['options']}")
    print(f"   Ans: {q['ans']}")
