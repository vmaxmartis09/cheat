# -*- coding: utf-8 -*-
import re
import openpyxl

wb = openpyxl.load_workbook('docs2/extracted/Session 7_ Correction of Part VII assignment/Session 7: Correction of Part VII assignment/master answer sheet for assignment part 7 v1.0.xlsx', data_only=True)
ws = wb['answer Key for test 1 (YBM)']
ybm_answers = {}
for r in range(1, 35):
    for c in range(1, 20):
        val = ws.cell(r, c).value
        if isinstance(val, int) and 101 <= val <= 200:
            ans = ws.cell(r, c+1).value
            if ans and str(ans).strip().upper() in ['A', 'B', 'C', 'D']:
                ybm_answers[val] = str(ans).strip().upper()

with open('docs2/extracted/ybm_part7_ocr.txt') as f:
    text = f.read()

text = re.sub(r'=== PAGE \d+ ===', '', text)
text = re.sub(r'YBM RC 1000.*', '', text)
text = re.sub(r'GO ON TO THE NEXT PAGE', '', text)

matches = list(re.finditer(r'Questions?\s+(\d+)\s*[-–]\s*(\d+)', text))

ybm_p7_list = []

for i, m in enumerate(matches):
    sq, eq = int(m.group(1)), int(m.group(2))
    start_pos = m.start()
    end_pos = matches[i+1].start() if i+1 < len(matches) else len(text)
    block = text[start_pos:end_pos]
    
    first_q_match = re.search(rf'\b{sq}\.\s+', block)
    passage_content = block[:first_q_match.start()].strip() if first_q_match else ""
    questions_block = block[first_q_match.start():].strip() if first_q_match else ""
    
    for q_num in range(sq, eq+1):
        if q_num < eq:
            next_q_pat = rf'\b{q_num+1}\.\s+'
            q_match = re.search(rf'\b{q_num}\.\s+(.*?)(?={next_q_pat})', questions_block, re.DOTALL)
        else:
            q_match = re.search(rf'\b{q_num}\.\s+(.*)', questions_block, re.DOTALL)
            
        if not q_match:
            continue
            
        q_raw = q_match.group(1).strip()
        opt_matches = list(re.finditer(r'\(([A-D])\)\s*([^(\n]+(?:\n(?!\([A-D]\))[^\n]+)*)', q_raw))
        q_text = q_raw[:opt_matches[0].start()].strip() if opt_matches else q_raw
        
        options = {}
        for om in opt_matches:
            letter = om.group(1)
            val = ' '.join(om.group(2).strip().split())
            options[letter] = val
            
        ans = ybm_answers.get(q_num, 'A')
        ybm_p7_list.append({
            'num': q_num,
            'question': q_text,
            'options': options,
            'ans': ans,
            'opt_count': len(options)
        })

print(f"Parsed {len(ybm_p7_list)} YBM Part 7 questions")
missing = [q['num'] for q in ybm_p7_list if q['opt_count'] < 4]
print("Missing opts:", missing)
for q in ybm_p7_list[:3]:
    print(f"Q{q['num']}: {q['question']}")
    print(f"   Opts: {q['options']}")
    print(f"   Ans: {q['ans']}")
